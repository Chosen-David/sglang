import sys
sys.path.insert(0, '/home/wangyuanshuo02/sglang/.local/lib/python3.11/site-packages')
from pptx import Presentation
import zipfile
import os

# 先解压zip
with zipfile.ZipFile('/home/wangyuanshuo02/sglang/two_level_indexer_research_pack.zip', 'r') as z:
    z.extractall('/home/wangyuanshuo02/sglang/indexer_proposal_extracted')

pptx_path = '/home/wangyuanshuo02/sglang/indexer_proposal_extracted/indexer_proposal/two_level_indexer_proposal_CN.pptx'
prs = Presentation(pptx_path)

print(f'共 {len(prs.slides)} 页幻灯片\n')
print('='*80)

for i, slide in enumerate(prs.slides, 1):
    print(f'\n--- 第 {i} 页 ---')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    print(text)
    # 检查备注
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print(f'[备注]: {notes}')

print('\n' + '='*80)
print(f'\n幻灯片宽度: {prs.slide_width}, 高度: {prs.slide_height}')
