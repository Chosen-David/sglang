# 提取 rework 版 PPT 全文（分析 GPT 批判意见与改动）
from pptx import Presentation

prs = Presentation('/home/wangyuanshuo02/sglang/TLI_progress_v4_critical_rework.pptx')
print(f'共 {len(prs.slides)} 页')
for i, slide in enumerate(prs.slides, 1):
    print(f'\n===== 第 {i} 页 =====')
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                print(t)
        if shape.shape_type == 13:
            print('[图片]')
    if slide.has_notes_slide:
        n = slide.notes_slide.notes_text_frame.text.strip()
        if n:
            print(f'--- 备注: {n}')
