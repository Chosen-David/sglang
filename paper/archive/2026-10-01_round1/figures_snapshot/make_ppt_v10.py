# TLI 终局数据 PPT v10（2026-09-30 10:00）
# v9 增量：①E72 mavg RULER 补全（87.59，主表双行格局定形）；②E67 tau02 完整梯度；
# ③新增降维页（E65/E66/E73 完整故事，用户指令）
# 数据来源：paper_main_table.json / e72_screen_verdict.json / b7s_ruler_final.json /
#           ruler_e72mavg.json / pred_e67_tau0{1,2}_score.json / e65/e73 JSON
import json

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

BLUE = RGBColor(0x2F, 0x6F, 0x9F)
RED = RGBColor(0xC1, 0x44, 0x3C)
GREEN = RGBColor(0x3A, 0x7D, 0x44)
GRAY = RGBColor(0x55, 0x55, 0x55)
DARK = RGBColor(0x22, 0x22, 0x22)

R = "/home/wangyuanshuo02/two-level-attention/exp/trace/results"
main_tbl = json.load(open(f"{R}/paper_main_table.json"))
screen = json.load(open(f"{R}/e72_screen_verdict.json"))
b7s_ruler = json.load(open("/tmp/tli_chain/b7s_ruler_final.json"))
e72m = json.load(open("/home/wangyuanshuo02/two-level-attention/exp/results_ruler/ruler_e72mavg.json"))
tau01 = json.load(open("/home/wangyuanshuo02/sglang/pred_e67_tau01_score.json"))
tau02 = json.load(open("/home/wangyuanshuo02/sglang/pred_e67_tau02_score.json"))
e73 = json.load(open(f"{R}/e73_nope_reduction.json"))

OUT = "/home/wangyuanshuo02/sglang/TLI_progress_v10.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def title_bar(s, text, sub=None, color=BLUE):
    from pptx.util import Inches as I
    box = s.shapes.add_textbox(I(0.5), I(0.25), I(12.3), I(0.85))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = color
    if sub:
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(12)
        p2.font.color.rgb = GRAY


def bullet_box(s, x, y, w, h, items, fs=14):
    from pptx.util import Inches as I
    box = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        color = DARK
        text = it
        if isinstance(it, tuple):
            text, color = it
        p.text = "• " + text if not text.startswith("   ") else text
        p.font.size = Pt(fs)
        p.font.color.rgb = color


def mono_table(s, x, y, w, h, rows, fs=12.5, highlight_col=None):
    from pptx.util import Inches as I
    tb = s.shapes.add_table(len(rows), len(rows[0]), I(x), I(y), I(w), I(h)).table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tb.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(fs if ri else fs + 0.5)
                p.font.bold = ri == 0
                if ri == 0:
                    p.font.color.rgb = DARK
                elif highlight_col is not None and ci == highlight_col:
                    p.font.color.rgb = GREEN
                    p.font.bold = True


def chip(s, x, y, w, h, big, small, color=GREEN):
    from pptx.util import Inches as I
    box = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = big
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = color
    p2 = tf.add_paragraph()
    p2.text = small
    p2.font.size = Pt(11)
    p2.font.color.rgb = GRAY


# ============ S1 封面：双臂格局终判 ============
s = slide()
title_bar(s, "TLI 终局判决（v10，2026-09-30）",
          "双臂格局定形：B7s 双基准最优行 + E72 LongBench 最优行；E67 gate 完整梯度；降维完整故事")
chip(s, 0.8, 1.6, 3.7, 1.2, "50.54 / 87.59",
     "TLI_E72（mavg β.375）\nLongBench 新高超 FullKV +0.18\nRULER 87.59 诚实并报", GREEN)
chip(s, 4.8, 1.6, 3.7, 1.2, "50.41 / 87.93",
     "TLI_B7s（β.25）双基准最优行\nLongBench 双超 + RULER 反超\nL16384 +0.91 长度梯度", GREEN)
chip(s, 8.8, 1.6, 3.7, 1.2, "34.76",
     "musique e2e F1 全方法新高\n（FullKV 32.14 +2.62）\n分区 far 检索机制证据", GREEN)
bullet_box(s, 0.8, 3.3, 12.0, 3.8, [
    ("双臂互补叙事：LongBench 多跳 far-heavy β.375 占优（musique +1.94）vs RULER 合成 needle β.25 占优（−0.34）——预算配比的任务形态依赖性", BLUE),
    ("口径鸿沟（方法论贡献 #8）：trace mass 排名与 e2e F1 排名五臂系统性反转——双口径必须都测都报告", RED),
    "E67 gate 完整梯度判决（τ0.005/0.015 双臂）：真触发区间微损不崩 → gate = 安全省算力保守开关",
    "降维完整故事（E65/E66/E73）：tail32 两段互补机理 + far 细筛 PCA d16 唯一推荐 + near 可压有兜底",
    "负结果资产：聚类代表 / 监督投影 e2e 崩 / TC 化反慢 / nope 全链 / 异 method 组合无增益",
], fs=13)

# ============ S2 主表 v2：双基准双臂 ============
s = slide()
title_bar(s, "论文主表 v2：双基准双臂终值",
          "paper_main_table.json；LongBench 13 任务 + RULER 33 任务（三长度平均）")
rows = [("method", "LongBench", "RULER", "定位")]
hl = {m["method"]: m for m in main_tbl["rows"]}
pos = {"TLI_B7": "双基准最优行", "TLI_E72": "LongBench 最优行"}
for m in ("FullKV", "Quest", "TIA", "TLI_C0", "TLI_B7", "TLI_E72"):
    d = hl[m]
    rows.append((m, str(d["longbench_avg"]),
                 str(d["ruler_avg"]) if d["ruler_avg"] is not None else "-",
                 pos.get(m, "")))
mono_table(s, 0.8, 1.45, 11.8, 3.0, rows, fs=14)
bullet_box(s, 0.8, 4.6, 12.0, 2.6, [
    ("TLI_E72 RULER 补全终值 87.59（91.34/88.33/83.08）：差距 vs B7s 全部集中于 multikey_3/cwe 两 4bit 粒度难任务族", GREEN),
    "TLI_B7 vs FullKV 长度梯度干净：−0.14（4K）→ −0.25（8K）→ +0.91（16K 反超）——超长上下文是分区 far 收益区",
    ("musique 34.76（FullKV +2.62）：多跳 far-heavy 任务独享大增益 = 预算分配增益最纯粹机制证据", GREEN),
    "Quest 双基准最弱（47.72/79.95）；TIA 同门消融基准（50.06/88.35）；vs Quest LongBench +2.8",
], fs=12.5)

# ============ S3 E72 五组合判决 ============
s = slide()
title_bar(s, "E72：五 method 组合 e2e 判决（快筛→全量兑现）",
          "e72_screen_verdict.json；trace mass 与 e2e F1 排名系统性反转")
rows = [("arm (far,near)", "α/β", "hotpotqa", "musique", "screen avg", "trace mass", "trace rank")]
order = ["mavg", "aavg", "mminmax", "mavg(B7s)", "cavg"]
tr = {"mavg": 4, "aavg": 5, "mminmax": 2, "mavg(B7s)": 1, "cavg": 3}
for a in order:
    d = screen[a]
    rows.append((a.replace("(B7s)", "\nβ=.25 ref"), f"{d['alpha']}/{d['beta']}",
                 str(d["hotpotqa"]), str(d["musique"]), str(d["screen_avg"]),
                 f"{d['trace_mass']:.4f}", str(tr[a])))
mono_table(s, 0.8, 1.45, 11.8, 3.2, rows, fs=11.5)
bullet_box(s, 0.8, 4.9, 12.0, 2.4, [
    ("全量终判：mavg 13 任务 AVG 50.54 = 系列新高（musique +1.94 幅度与快筛精确一致）", GREEN),
    ("trace 第 4 的 mavg → e2e 全量冠军；trace 垫底 aavg → e2e 第 2——「trace 排名不可替代 e2e 判决」获全量确认", RED),
    "β=0.375 > β=0.25（LongBench）：近端预算份额宜超区域份额；但 RULER 上反向（β.25 占优 −0.34）——任务形态依赖",
], fs=12.5)

# ============ S4 降维完整故事（E65/E66/E73，新增页） ============
s = slide()
title_bar(s, "降维与投影：完整消融故事（E65 + E66 + E73）",
          "e65_dim_reduction.json / E66 LOTO / e73_nope_reduction.json；trace 重放 8 样本×8 层")
mono_table(s, 0.8, 1.45, 11.8, 3.0, [
    ("配置", "far mass 捕获", "结论"),
    ("tail32（rope 低频尾16 + nope 尾16）", "0.796", "主表口径；两段互补 +0.22 组合增益"),
    ("rope 低频尾 16 单独", "0.576", "位置感知段（长程检索主力）"),
    ("nope 尾 16 单独", "0.383", "内容稳定段"),
    ("nope 全链（64 维最佳 trunc d32）", "0.435", "E73：差 36pt——纯 nope 否定"),
    ("细筛 PCA d16（无监督）", "0.766 vs 0.804", "E65：唯一推荐落地降维（近无损）"),
    ("训练投影 wsvd（E66 LOTO）", "基线 92-94%", "跨数据集成立但 e2e 崩——口径鸿沟 negative"),
], fs=11.5)
bullet_box(s, 0.8, 4.7, 12.0, 2.5, [
    ("tail32 互补性机理（E73 新发现）：32 维子空间的威力 = 位置感知段 + 内容稳定段各取所长——解释 random32 崩 vs tail32 稳", GREEN),
    "near 侧压维极平坦（d4≈d64）：信息冗余大，8 维即够 + 真实管线滑窗兜底",
    "MLA q/k 下投影：Qwen3 纯 GQA 满秩无自带低秩结构——SVD 近似即 E66 已判 negative；DeepSeek MLA latent 索引 = future work",
], fs=12.5)

# ============ S5 E67 gate 完整梯度 ============
s = slide()
title_bar(s, "E67：感知型 per-request gate 完整梯度判决",
          "pred_e67_tau0{1,2}_score.json；对照 = gate_off（musique 27.57 / qasper 40.37）")
mono_table(s, 0.8, 1.45, 11.8, 2.4, [
    ("τ", "触发程度", "musique", "qasper", "合计 Δ", "判决"),
    ("0.01（默认）", "近 no-op（1.1% 输出分歧）", "—", "—", "—", "「无损」证据强度不足"),
    ("0.005", "实质非零（11.8% 实体分歧）", "27.80 (+0.23)", "39.96 (−0.41)", "−0.18", "微损不崩"),
    ("0.015", "安全/多跳边界", "27.84 (+0.27)", "39.70 (−0.67)", "−0.40", "微损不崩"),
], fs=12.5)
bullet_box(s, 0.8, 4.2, 12.0, 2.9, [
    ("梯度单调（τ×3 → 损失×2.2）；qasper（far 总量 0.55-0.95）主要承受方，musique 反微升", GREEN),
    ("对照静态版掉 4.8-5.9 分：per-request 动态信号（prefill last1 corr 0.924）防住反向错误", GREEN),
    ("gate 终定位 = 「安全省算力的保守开关」而非精度增益点：正收益未兑现", BLUE),
    "速度侧资产：跳 near 空间 31% 层/2.3% mass@τ0.1 > 跳 far 12%/1.7%；τ≥0.2 才有可观跳层空间",
], fs=12.5)

# ============ S6 消融定稿 ============
s = slide()
title_bar(s, "消融定稿：六维参数空间（E64a-j，132 臂 + E72 五臂 e2e）",
          "用户一般化框架：mid 分 near/far → 降维 → 双区各自粗筛+细筛")
mono_table(s, 0.8, 1.45, 11.8, 3.0, [
    ("参数", "终值", "依据"),
    ("α（near 区比例）", "0.125", "51+81 臂网格单调递减；最优内部臂恒 α.125"),
    ("β（near 页预算）", "0.375 / 0.25 双臂", "任务形态依赖：LongBench 多跳 β.375 占优 / RULER 合成 needle β.25 占优（双臂并报）"),
    ("γ（near token 预算）", "0.125", "必须随 K2 等比缩放；γ.5 在 K2=1024 far 崩（B0 13.51 教训）"),
    ("far_method", "minmax", "块上界 mono 最稳 0.9166；e2e 冠军臂"),
    ("near_method", "avg", "near 侧自由度 ≤0.003（E64i 双预算口径）；异 method 组合删"),
    ("降维", "tail32 / 细筛 PCA d16", "tail32 两段互补机理（E73）；监督投影 e2e 崩（negative）"),
], fs=11.5)
bullet_box(s, 0.8, 4.7, 12.0, 2.5, [
    ("分区价值主张：near 免 min/max 打分 = 结构性带宽节省 + 紧预算分配质量（e2e 双口径判决）", GREEN),
    "kernel 资产：fused L1 3.6× / L2 级联 1.63× / M11 prefill e2e 4.30×；「L1 容多选、L2 不能」级联边界",
    ("negative results 全家桶：聚类代表 / TC 化反慢 0.63× / 在线信号失效 / 跨层复用失效 / nope 全链", RED),
], fs=12.5)

# ============ S7 结论页 ============
s = slide()
title_bar(s, "结论与剩余工作",
          "全部主线实验收官；#74 创新点重评估 + paper 重写进行中")
bullet_box(s, 0.8, 1.5, 12.0, 5.5, [
    ("主结果：LongBench 50.54（E72）/ RULER 87.93（B7s）——双臂并报，training-free 稀疏索引器双基准逼近或超越 FullKV", GREEN),
    "创新点定稿：A 子空间粗筛（tail32 两段互补机理）/ B' 分区预算（e2e 成立）/ D' gate（安全开关消融）/ 口径鸿沟（方法论贡献）",
    "速度三支柱：kernel fused 1.63-3.6× / 30B 64K e2e 1.285× / 8B 稳态 decode 1.21×——对拍精确一致 + 逐字一致",
    "双口径铁律贯穿：kernel microbench 与 e2e 都测都报；trace 排名不可替代 e2e 判决（五臂全反转实证）",
    ("剩余：paper 全文重写（tex v2 已更新至终局数字）+ 投稿打磨", BLUE),
], fs=14)

prs.save(OUT)
print("saved ->", OUT)
